"""Unit tests for the pure, I/O-free parts of the connector."""

from __future__ import annotations

from datetime import datetime, timedelta, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

from now_mcp.extract import extract_text, html_to_text, snippets
from now_mcp.service import (
    build_deadlines,
    flatten_toc,
    is_current_module,
    match_status_row,
    shape_assignment,
    shape_module,
)
from now_mcp.util import days_until, fmt_local, parse_iso, safe_name

TZ = ZoneInfo("Europe/London")
NOW = datetime(2026, 10, 9, 20, 0, tzinfo=timezone.utc)


def test_parse_iso_handles_brightspace_format():
    dt = parse_iso("2026-10-30T12:00:00.000Z")
    assert dt == datetime(2026, 10, 30, 12, 0, tzinfo=timezone.utc)
    assert parse_iso(None) is None
    assert parse_iso("garbage") is None


def test_fmt_local_uses_timezone():
    assert fmt_local("2026-10-30T12:00:00.000Z", TZ) == "Fri 30 Oct 2026, 12:00"
    # BST: 1 hour ahead of UTC in early October.
    assert fmt_local("2026-10-09T12:00:00.000Z", TZ) == "Fri 09 Oct 2026, 13:00"


def test_days_until():
    assert days_until("2026-10-11T20:00:00.000Z", NOW) == 2.0
    assert days_until("2026-10-08T20:00:00.000Z", NOW) == -1.0


def test_safe_name():
    assert safe_name("SOFT30111: Software/Eng <2026>") == "SOFT30111_ Software_Eng _2026"
    assert safe_name("   ") == "untitled"


def test_shape_module_and_current_filter():
    enrollment = {
        "OrgUnit": {"Id": 123, "Code": "SOFT30111", "Name": "Software Engineering", "Type": {"Name": "Course Offering"}},
        "Access": {"IsActive": True, "CanAccess": True, "StartDate": "2026-09-20T00:00:00.000Z", "EndDate": "2027-06-30T00:00:00.000Z"},
        "PinDate": None,
    }
    mod = shape_module(enrollment)
    assert mod["ou"] == 123 and mod["code"] == "SOFT30111"
    assert is_current_module(mod, NOW)
    old = dict(mod, end="2025-06-30T00:00:00.000Z")
    assert not is_current_module(old, NOW)
    inactive = dict(mod, is_active=False)
    assert not is_current_module(inactive, NOW)


def test_flatten_toc_builds_paths():
    toc = {
        "Modules": [
            {
                "ModuleId": 1,
                "Title": "Week 1",
                "Topics": [{"TopicId": 10, "Title": "Lecture 1", "TypeIdentifier": "File", "Url": "/content/enforced/x/lec1.pdf", "Unread": True}],
                "Modules": [
                    {
                        "ModuleId": 2,
                        "Title": "Lab",
                        "Topics": [{"TopicId": 11, "Title": "Panopto", "TypeIdentifier": "Link", "Url": "https://panopto.example/x"}],
                        "Modules": [],
                    }
                ],
            }
        ]
    }
    topics = flatten_toc(toc)
    assert [t["topic_id"] for t in topics] == [10, 11]
    assert topics[0]["path"] == "Week 1" and topics[0]["downloadable"] is True and topics[0]["unread"] is True
    assert topics[1]["path"] == "Week 1 / Lab" and topics[1]["downloadable"] is False


def test_shape_assignment_fields():
    folder = {
        "Id": 55,
        "Name": "Coursework 1: Design Report",
        "DueDate": "2026-11-06T13:00:00.000Z",
        "Availability": {"StartDate": "2026-09-30T00:00:00.000Z", "EndDate": "2026-11-13T13:00:00.000Z"},
        "GroupTypeId": 7,
        "Assessment": {"ScoreDenominator": 100, "Rubrics": [{"RubricId": 1}]},
        "CustomInstructions": {"Text": "Submit a PDF.", "Html": "<p>Submit a PDF.</p>"},
        "Attachments": [{"FileId": 9, "FileName": "brief.pdf", "Size": 1234}],
    }
    mod = {"ou": 123, "code": "SOFT30111", "name": "Software Engineering"}
    a = shape_assignment(folder, mod, "https://now.ntu.ac.uk", TZ)
    assert a["is_group"] is True and a["marks"] == 100 and a["has_rubric"] is True
    assert a["due_local"] == "Fri 06 Nov 2026, 13:00"
    assert a["attachments"][0]["name"] == "brief.pdf"
    assert "db=55" in a["link"] and "ou=123" in a["link"]


def test_match_status_row():
    tables = [[
        ["Assignment", "Completion Status", "Score", "Evaluation Status", "Due Date"],
        ["Coursework 1: Design Report Due on 6 Nov 2026 1:00 PM", "Not Submitted", "- / 100", "Not yet evaluated", "6 Nov 2026"],
        ["Coursework 2 Due on 1 Dec 2026", "1 Submission, 1 File Submitted", "- / 100", "", ""],
    ]]
    assert match_status_row("Coursework 1: Design Report", tables)["submitted"] is False
    assert match_status_row("Coursework 2", tables)["submitted"] is True
    assert match_status_row("Exam", tables) is None


def test_build_deadlines_merges_sorts_and_dedupes():
    assignments = [
        {"type": "assignment", "title": "Coursework 1", "due": "2026-10-20T13:00:00.000Z"},
        {"type": "assignment", "title": "Old thing", "due": "2026-10-01T13:00:00.000Z"},  # 8 days overdue
        {"type": "assignment", "title": "Far future", "due": "2027-05-01T13:00:00.000Z"},  # outside window
    ]
    quizzes = [{"type": "quiz", "title": "Quiz 1", "due": "2026-10-15T23:59:00.000Z"}]
    events = [
        {"type": "event", "title": "Coursework 1", "start": "2026-10-20T13:00:00.000Z"},  # duplicate of the assignment
        {"type": "event", "title": "Exam", "start": "2026-12-01T09:00:00.000Z"},
    ]
    items = build_deadlines(assignments, quizzes, events, days_back=14, days_ahead=60, now=NOW)
    titles = [i["title"] for i in items]
    assert titles == ["Old thing", "Quiz 1", "Coursework 1", "Exam"]
    assert items[0]["overdue"] is True and items[0]["days_left"] < 0
    assert items[2]["type"] == "assignment"  # assignment wins over the duplicate event


def test_html_to_text_and_snippets():
    text = html_to_text("<html><body><script>x()</script><h1>Title</h1><p>Hello   world</p></body></html>")
    assert "x()" not in text and "Title" in text and "Hello world" in text
    assert snippets("the quick brown fox jumps over the lazy dog", "brown", width=10) == ["uick brown fox"]


def test_extract_text_pdf_docx_pptx(tmp_path: Path):
    from docx import Document
    from pptx import Presentation
    from pypdf import PdfWriter

    d = Document()
    d.add_paragraph("Agile delivery sprint")
    d.save(tmp_path / "a.docx")
    text, err = extract_text(tmp_path / "a.docx")
    assert err is None and "Agile delivery sprint" in text

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Lecture 3: Testing"
    prs.save(tmp_path / "b.pptx")
    text, err = extract_text(tmp_path / "b.pptx")
    assert err is None and "Lecture 3: Testing" in text and "--- slide 1 ---" in text

    w = PdfWriter()
    w.add_blank_page(width=200, height=200)
    with open(tmp_path / "c.pdf", "wb") as fh:
        w.write(fh)
    text, err = extract_text(tmp_path / "c.pdf")
    assert err is None and "--- page 1 ---" in text

    (tmp_path / "d.mp4").write_bytes(b"\x00\x01")
    text, err = extract_text(tmp_path / "d.mp4")
    assert text == "" and err


def test_new_content_window_and_seen(tmp_path: Path, monkeypatch):
    """new_content: window filter, seen-set and mark_seen, using a fake client."""
    import asyncio
    from now_mcp.config import Config
    from now_mcp.service import NowService

    cfg = Config(state_dir=tmp_path / "s", data_dir=tmp_path / "d")
    cfg.state_dir.mkdir(); cfg.data_dir.mkdir()
    fresh = (NOW - timedelta(hours=5)).strftime("%Y-%m-%dT%H:%M:%S.000Z")
    stale = (NOW - timedelta(days=10)).strftime("%Y-%m-%dT%H:%M:%S.000Z")

    class FakeClient:
        async def toc(self, ou):
            return {"Modules": [{"ModuleId": 1, "Title": "Week 5", "Topics": [
                {"TopicId": 1, "Title": "New slides", "TypeIdentifier": "File", "LastModifiedDate": fresh},
                {"TopicId": 2, "Title": "Old slides", "TypeIdentifier": "File", "LastModifiedDate": stale},
            ], "Modules": []}]}
        async def news(self, ou):
            return [{"Id": 9, "Title": "Room change", "StartDate": fresh, "Body": {"Text": "x", "Html": "x"}}]
        async def dropbox_folders(self, ou):
            return [{"Id": 7, "Name": "CW2", "DueDate": fresh}]

    svc = NowService(cfg, FakeClient())
    svc._modules = {5: {"ou": 5, "code": "X", "name": "X", "is_active": True, "can_access": True, "pinned": False}}
    svc._modules_at = 10**12

    import now_mcp.service as S
    monkeypatch.setattr(S, "utcnow", lambda: NOW)  # freeze time for the window
    r1 = asyncio.run(svc.new_content(hours=72, mark_seen=True))
    assert [t["title"] for t in r1["topics"]] == ["New slides"]
    assert r1["counts"] == {"topics": 1, "announcements": 1, "assignments": 1}
    r2 = asyncio.run(svc.new_content(hours=72))
    assert r2["counts"] == {"topics": 0, "announcements": 0, "assignments": 0}
    r3 = asyncio.run(svc.new_content(hours=72, only_unseen=False))
    assert r3["counts"]["topics"] == 1


def test_login_flow_target_url_is_restricted_to_ntu(tmp_path: Path):
    from now_mcp.config import Config
    from now_mcp.session import LoginFlow

    cfg = Config(state_dir=tmp_path / "s", data_dir=tmp_path / "d")
    f = LoginFlow(cfg, target_url="https://myntuac.sharepoint.com/sites/DT")
    assert f.target_host == "myntuac.sharepoint.com"
    assert LoginFlow(cfg, target_url="https://services.ntu.ac.uk/Timetabling/").target_host == "services.ntu.ac.uk"
    assert LoginFlow(cfg).target_url is None
    import pytest
    with pytest.raises(ValueError):
        LoginFlow(cfg, target_url="https://evil.example.com/")
